from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference, PieChart

def create_report(orders1, orders2, name):
    wb = Workbook()
    wb.create_sheet(title='Первый лист', index=0)
    
    sheet1 = wb['Первый лист']
    sheet1['A1'] = 'Автор'
    sheet1['A2'] = 'Количество'
    
    idx = 1
   
    for order in orders1:
        cell = sheet1.cell(row=1, column=idx+1)
        cell.value = order.author.last_name

        cell = sheet1.cell(row=2, column=idx+1)
        cell.value = order.count

        idx += 1


    chart1 = BarChart()
    chart1.title = 'Соотношение пользователей по числу отправленных заявок'
    data1 = Reference(sheet1, min_col=2, min_row=1, max_col=len(orders1) + 1, max_row=2)
    chart1.add_data(data1, titles_from_data=True)
    sheet1.add_chart(chart1, 'A4')

    wb.create_sheet(title='Второй лист', index=1)

    sheet2 = wb['Второй лист']
    sheet2['A1'] = 'Статус'
    sheet2['A2'] = 'Количество'

    idx = 1

    for order in orders2:
        cell = sheet2.cell(row=1, column=idx+1)
        cell.value = order.status

        cell = sheet2.cell(row=2, column=idx+1)
        cell.value = order.count

        idx += 1
    
    chart2 = PieChart()
    chart2.title = 'Отчет о выполняемости заявок'
    labels = Reference(sheet2, min_col=2, min_row=1, max_col=len(orders2) + 1, max_row=1)
    data2 = Reference(sheet2, min_col=2, min_row=2, max_col=len(orders2) + 1, max_row=2)
    chart2.add_data(data2, titles_from_data=True)
    chart2.set_categories(labels)
    sheet2.add_chart(chart2, 'A4')

    wb.save(f'reports/{name}.xlsx')


    